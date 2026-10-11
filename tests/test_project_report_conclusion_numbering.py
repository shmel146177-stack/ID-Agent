from docx import Document

from app.generators.project_report_generator import ProjectReportGenerator


def test_project_report_conclusion_is_section_seven():
    generator = ProjectReportGenerator()
    document = Document()

    generator._add_conclusion(
        document,
        {},
    )

    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "7. Заключение ID-Agent" in full_text


def test_project_report_does_not_call_matched_positions_complete_sheets():
    generator = ProjectReportGenerator()
    document = Document()
    generator._add_conclusion(
        document,
        {
            "determined": True,
            "required_count": 22,
            "found_count": 22,
            "missing_count": 0,
            "completeness_percent": None,
            "position_completeness_percent": 100.0,
            "analysis_coverage": {
                "status": "incomplete",
                "documents": [
                    {"filename": "Том 1.pdf", "missing_pages": [27, 36]},
                    {"filename": "Том 2.pdf", "missing_pages": [12, 18, 20]},
                ],
            },
        },
    )

    conclusion = document.paragraphs[1].text
    assert "22 позиций" in conclusion
    assert "Общая комплектность автоматически не определена" in conclusion
    assert "Том 1.pdf: 27, 36" in conclusion
    assert "Том 2.pdf: 12, 18, 20" in conclusion
    assert "полный" not in conclusion
    assert "None%" not in conclusion


def test_equipment_conclusion_counts_document_categories():
    document = Document()
    ProjectReportGenerator()._add_conclusion(
        document,
        {
            "profile": "equipment",
            "required_count": 3,
            "found_count": 2,
            "missing_count": 1,
            "completeness_percent": 66.7,
        },
    )

    assert "2" in document.paragraphs[1].text
    assert "категорий документов" in document.paragraphs[1].text
    assert "позиций ведомости" not in document.paragraphs[1].text
