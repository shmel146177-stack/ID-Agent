from openpyxl import load_workbook

import app.generators.document_registry_excel as excel_module
from app.generators.document_registry_excel import DocumentRegistryExcel


def test_document_registry_excel_creates_real_xlsx(monkeypatch, tmp_path):

    generator = DocumentRegistryExcel()

    registry = {
        "documents": [
            {
                "number": 1,
                "filename": "project.pdf",
                "classification": "Чертеж",
                "status": "Обработан",
                "extension": ".pdf",
                "drawing_number": "TEST-001",
                "date": "12.08.2026",
                "manufacturer": "ООО Тест",
                "equipment": "Тестовое оборудование",
            }
        ]
    }

    completeness = {
        "status": "Полный комплект",
        "profile": "project",
        "profile_name": "Проектная документация",
        "check_method": "Тестовая проверка",
        "required_count": 1,
        "found_count": 1,
        "missing_count": 0,
        "completeness_percent": 100.0,
        "documents": [
            {
                "sheet_number": 1,
                "title": "Общие данные",
                "status": "Есть",
                "matched_page": 1,
                "matched_filename": "project.pdf",
                "matched_page_type": "Общие данные",
                "score": 120,
                "confidence": "Высокая",
            }
        ],
        "missing_sheets": [],
    }

    monkeypatch.setattr(
        excel_module.document_registry,
        "build",
        lambda project_name: registry,
    )

    monkeypatch.setattr(
        excel_module.document_completeness,
        "check",
        lambda project_name: completeness,
    )

    output_path = tmp_path / "document_registry.xlsx"

    monkeypatch.setattr(
        generator,
        "_output_path",
        lambda project_name: output_path,
    )

    result = generator.create("TEST_PROJECT")

    assert result == str(output_path)
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    workbook = load_workbook(output_path)

    assert len(workbook.worksheets) == 2

    registry_sheet = workbook.worksheets[0]
    completeness_sheet = workbook.worksheets[1]

    assert "TEST_PROJECT" in str(registry_sheet["A2"].value)
    assert registry_sheet["B5"].value == "project.pdf"
    assert registry_sheet["C5"].value == "Чертеж"
    assert registry_sheet["F5"].value == "TEST-001"

    assert "TEST_PROJECT" in str(completeness_sheet["A2"].value)

    assert completeness_sheet["B6"].value == 1
    assert completeness_sheet["B7"].value == 1
    assert completeness_sheet["B8"].value == 0
    assert completeness_sheet["B9"].value == "100.0%"
    assert completeness_sheet["A6"].value == "Позиций в ведомости"
    assert completeness_sheet["A7"].value == "Сопоставлено позиций"
    assert completeness_sheet["A8"].value == "Не сопоставлено позиций"

    workbook.close()


def test_document_registry_excel_shows_attachment_warning(monkeypatch, tmp_path):
    generator = DocumentRegistryExcel()
    monkeypatch.setattr(
        excel_module.document_registry, "build", lambda name: {"documents": []}
    )
    monkeypatch.setattr(
        excel_module.document_completeness,
        "check",
        lambda name: {
            "status": "Требует проверки: расхождения в прилагаемых документах",
            "profile": "project",
            "required_count": 1,
            "found_count": 1,
            "missing_count": 0,
            "completeness_percent": None,
            "documents": [],
            "missing_sheets": [],
            "attachment_audit": {
                "warnings": [
                    {
                        "message": "Ведомость: 98765/26-КЖ.СП2; "
                        "кандидат в штампе: 98766/26-КЖ.СП2 (PDF 4)"
                    }
                ]
            },
        },
    )
    output_path = tmp_path / "registry.xlsx"
    monkeypatch.setattr(generator, "_output_path", lambda name: output_path)

    generator.create("OTHER_PROJECT")

    workbook = load_workbook(output_path)
    sheet = workbook["Комплектность"]
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
    assert "Прилагаемые документы — требуется проверка" in values
    assert any("98766/26-КЖ.СП2 (PDF 4)" in str(value) for value in values)
    assert sheet["B5"].fill.fgColor.rgb == "00FFF2CC"
    workbook.close()
