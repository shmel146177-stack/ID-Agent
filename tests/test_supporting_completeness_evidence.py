from docx import Document

from app.generators.project_document_set import ProjectDocumentSet
from app.generators.project_report_generator import ProjectReportGenerator
from app.services.project_package import ProjectPackage
from app.services.document_completeness import DocumentCompleteness
import app.services.document_completeness as completeness_module


def test_requirement_evidence_separates_match_candidate_and_missing(
    monkeypatch, tmp_path
):
    document_set = ProjectDocumentSet()
    monkeypatch.setattr(document_set, "_analysis_path", lambda name: tmp_path)

    def load_analysis(path):
        if path.name == "project_analysis.json":
            return {
                "documents": [
                    {
                        "filename": "cable.pdf",
                        "classification": "Схема",
                        "status": "Обработан",
                    },
                    {
                        "filename": "unknown.pdf",
                        "classification": "Схема",
                        "status": "Ошибка",
                    },
                ]
            }
        return {"documents": [{"filename": "cable.pdf", "pages": [{"text": "кабель"}]}]}

    monkeypatch.setattr(document_set, "_load_json", load_analysis)
    requirements = [
        {
            "code": "cable",
            "title": "Кабельная схема",
            "reason": "Акт требует схему",
            "source_act": {"code": "act-1"},
            "evidence": [{"sheet_number": "2"}],
            "document_types": ["Схема"],
            "match_keywords": ["кабель"],
        },
        {"code": "ground", "title": "Заземление", "match_keywords": ["заземление"]},
        {"code": "test", "title": "Испытание", "match_keywords": ["испытание"]},
    ]
    files = [{"name": "cable.pdf", "path": "04/cable.pdf"}, {"name": "unknown.pdf"}]
    result = document_set._supporting_section_completeness(
        "TEST", {"required_count": 3, "documents": requirements}, files
    )

    assert (result["found_count"], result["review_count"], result["missing_count"]) == (
        1,
        1,
        1,
    )
    matched, pending_ground, pending_test = result["requirement_assessments"]
    assert matched["status"] == "Совпадение по анализу"
    assert matched["basis"]["matched_file"] == files[0]
    assert matched["basis"]["source_act"] == {"code": "act-1"}
    assert matched["basis"]["source_evidence"] == [{"sheet_number": "2"}]
    assert matched["basis"]["match_rule"] == {
        "document_types": ["Схема"],
        "match_keywords": ["кабель"],
    }
    for pending in (pending_ground, pending_test):
        assert pending["status"] == "Не подтверждено; требуется проверка"
        assert pending["basis"]["review_candidates"] == [files[1]]
        assert "matched_file" not in pending["basis"]
        assert pending["engineer_confirmation_required"] is True
    assert result["count_basis"]["minimum_missing_requirements"] == 1


def test_report_and_manifest_show_requirement_basis(monkeypatch, tmp_path):
    report = ProjectReportGenerator()
    monkeypatch.setattr(report, "_project_path", lambda name: tmp_path)
    monkeypatch.setattr(
        ProjectDocumentSet,
        "_list_section_files",
        lambda self, path: [{"name": "unknown.pdf"}],
    )
    section = {
        "number": "04",
        "code": "executive_schemes",
        "title": "Исполнительные схемы",
        "required_count": 1,
        "documents": [
            {"code": "scheme", "title": "Схема", "reason": "Требование акта"}
        ],
    }
    enriched = report._with_supporting_completeness("TEST", {"sections": [section]})
    assessment = enriched["sections"][0]["requirement_assessments"][0]
    assert assessment["basis"]["requirement_reason"] == "Требование акта"

    document = Document()
    report._add_supporting_documents(document, enriched)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "основание требования — Требование акта" in text
    assert "файлов без полного анализа — 1" in text

    manifest_section = {
        "number": "04",
        "code": "executive_schemes",
        "title": "Исполнительные схемы",
        "status": "Требует проверки",
        "path": tmp_path,
        "detected": enriched["sections"][0],
    }
    manifest = ProjectPackage()._build_document_sections(
        {"sections": [manifest_section]}, tmp_path
    )[0]
    assert (
        manifest["requirement_assessments"]
        == enriched["sections"][0]["requirement_assessments"]
    )
    assert manifest["count_basis"]["unanalysed_files"] == 1


def test_drawing_conclusions_keep_register_and_page_evidence(monkeypatch):
    monkeypatch.setattr(
        completeness_module.drawing_sheet_matcher,
        "analyze_project",
        lambda name: {
            "expected_count": 2,
            "found_count": 1,
            "missing_count": 1,
            "determined": True,
            "completeness_percent": 50.0,
            "matches": [
                {
                    "sheet_number": 1,
                    "title": "Схема",
                    "found": True,
                    "register_filename": "volume.pdf",
                    "matched_filename": "volume.pdf",
                    "matched_page": 3,
                    "matched_phrases": ["Схема"],
                    "score": 120,
                },
                {
                    "sheet_number": 2,
                    "title": "План",
                    "found": False,
                    "register_filename": "volume.pdf",
                    "score": 0,
                },
            ],
        },
    )
    result = DocumentCompleteness().check("TEST", profile="project")
    found, missing = result["documents"]
    assert found["basis"]["register_file"] == "volume.pdf"
    assert found["basis"]["matched_page"] == 3
    assert found["basis"]["matched_phrases"] == ["Схема"]
    assert missing["basis"]["matched_file"] is None
    assert missing["basis"]["register_sheet"] == 2

    document = Document()
    ProjectReportGenerator()._add_completeness(document, result)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "PDF — volume.pdf, страница 3" in text
    assert "PDF — не найден" in text


def test_equipment_conclusions_name_classified_files():
    completeness = DocumentCompleteness()
    completeness._get_registry = lambda name: {
        "documents": [
            {"filename": "passport.pdf", "classification": "Паспорт оборудования"}
        ]
    }
    result = completeness.check("TEST", profile="equipment")
    assert result["documents"][0]["basis"]["matched_files"] == ["passport.pdf"]
    assert result["documents"][1]["basis"]["matched_files"] == []


def test_unverifiable_requirement_keeps_unassigned_analysed_file_for_review(
    monkeypatch, tmp_path
):
    document_set = ProjectDocumentSet()
    monkeypatch.setattr(document_set, "_analysis_path", lambda name: tmp_path)
    monkeypatch.setattr(
        document_set,
        "_load_json",
        lambda path: {
            "documents": (
                [{"filename": "unknown.pdf", "classification": "Схема"}]
                if path.name == "project_analysis.json"
                else [{"filename": "unknown.pdf", "pages": [{"text": "прочее"}]}]
            )
        },
    )
    result = document_set._supporting_section_completeness(
        "TEST",
        {"required_count": 1, "documents": [{"code": "manual"}]},
        [{"name": "unknown.pdf"}],
    )
    assert result["review_count"] == 1
    assert result["count_basis"]["unanalysed_files"] == 0
    assert result["requirement_assessments"][0]["status"] == (
        "Не подтверждено; требуется проверка"
    )
